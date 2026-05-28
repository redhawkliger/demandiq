"""
DemandIQ — Export for Fine-tuning / feedback pipeline
=====================================================
Converts confirmed rows from labels_template.xlsx into JSONL reference
examples (the same shape count_agent.load_examples reads).

ONLY rows with confirmed == TRUE are exported. Wrong/unconfirmed
predictions must never become reference examples — that would poison the
in-context learning.

Two modes:
  (default)            Rewrite <output.jsonl> from all confirmed rows.
  --append-confirmed   Append ONLY confirmed rows that aren't already in
                       <output.jsonl> (dedupe by file_name + intersection).
                       Idempotent: running twice adds each example once.

Usage:
    python3 export_for_finetuning.py [<labels.xlsx>] [<output.jsonl>] \
        [--append-confirmed] [--files-root <dir>]
"""

import sys, json, re, argparse
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Run: pip install openpyxl --break-system-packages")

from count_agent import read_file, SYSTEM_PROMPT

MOVEMENTS = [
    'EBU','EBL','EBT','EBR',
    'WBU','WBL','WBT','WBR',
    'NBU','NBL','NBT','NBR',
    'SBU','SBL','SBT','SBR',
]

_FILENAME_RE = re.compile(r'Filename:\s*(.+)')


def dedupe_key(file_name, intersection):
    """Stable key for a training example: source file + intersection name.
    Used to avoid appending the same confirmed row twice."""
    fn = (file_name or '').strip()
    iv = (intersection or '').strip()
    return f"{fn}||{iv}"


def existing_keys(output_path):
    """Read the dedupe keys already present in an output JSONL. Returns a
    set; empty if the file doesn't exist yet. Parses the filename from each
    entry's user message ('Filename: X') and the intersection name from the
    assistant JSON — matching how build_training_example writes them."""
    keys = set()
    p = Path(output_path)
    if not p.exists():
        return keys
    with open(p, encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                ex = json.loads(line)
            except json.JSONDecodeError:
                continue
            msgs = ex.get('messages', [])
            user = next((m.get('content', '') for m in msgs if m.get('role') == 'user'), '')
            asst = next((m.get('content', '') for m in msgs if m.get('role') == 'assistant'), '')
            fname = None
            if isinstance(user, str):
                m = _FILENAME_RE.search(user)
                if m:
                    fname = m.group(1).strip()
            intersection = None
            try:
                d = json.loads(asst)
                ivs = d.get('intersections', [])
                if ivs:
                    intersection = ivs[0].get('name')
            except (json.JSONDecodeError, TypeError):
                pass
            keys.add(dedupe_key(fname, intersection))
    return keys


def row_to_extraction(row):
    """Convert an Excel row back into the extraction JSON structure."""
    volumes = {'AM': {}, 'PM': {}}
    for period in ('AM', 'PM'):
        for mv in MOVEMENTS:
            v = row.get(f'{mv}_{period}')
            volumes[period][mv] = v if v is not None else None

    return {
        "vendor_detected": row.get('vendor'),
        "intersections": [{
            "name":           row.get('intersection'),
            "am_peak_start":  row.get('AM_peak'),
            "pm_peak_start":  row.get('PM_peak'),
            "volumes":        volumes,
            "phf":            {"AM": row.get('PHF_AM'), "PM": row.get('PHF_PM')},
            "hv_pct":         {"AM": row.get('HV_AM_pct'), "PM": row.get('HV_PM_pct')},
            "hv_method":      row.get('HV_method', ''),
            "notes":          row.get('notes', '') or '',
        }]
    }


def build_training_example(file_content, filename, extraction):
    """Reference-example format — same shape count_agent.load_examples reads
    (system + user + assistant messages)."""
    return {
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user",   "content": f"Filename: {filename}\n\nFile content:\n{file_content}"},
            {"role": "assistant", "content": json.dumps(extraction, indent=None)}
        ]
    }


def read_confirmed_rows(labels_path):
    """Return list of row dicts where confirmed is True. Confirmed-only
    gating lives here — the single chokepoint that keeps unconfirmed
    predictions out of the training set."""
    wb = load_workbook(labels_path)
    ws = wb['Labels']

    header = {}
    for ci in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=ci).value
        if h:
            header[h] = ci
    if 'file_name' not in header or 'confirmed' not in header:
        sys.exit("Labels file missing required columns (file_name, confirmed)")

    rows = []
    for r in range(2, ws.max_row + 1):
        if not ws.cell(row=r, column=header['file_name']).value:
            continue
        if ws.cell(row=r, column=header['confirmed']).value is not True:
            continue   # GATE: only confirmed == True
        rows.append({h: ws.cell(row=r, column=ci).value for h, ci in header.items()})
    return rows


def _load_file_content(files_root, fname):
    """Resolve + read a source count file's content for the training entry.
    Returns the content string, or None if not found/unreadable."""
    if not files_root:
        return "[FILE CONTENT NOT LOADED]"
    fp = files_root / fname
    if not fp.exists():
        found = list(files_root.rglob(fname))
        if not found:
            return None
        fp = found[0]
    try:
        # read_file returns a 3-tuple (content, ftype, warnings) — Stage 7
        # fix: previously unpacked as 2 and crashed.
        content, _ftype, _warnings = read_file(str(fp))
        return content
    except Exception as e:
        print(f"  ✗ Read failed for {fname}: {e}")
        return None


def export(labels_path, output_path, files_root=None, append=False):
    confirmed_rows = read_confirmed_rows(labels_path)
    if not confirmed_rows:
        sys.exit("No confirmed rows found. Set confirmed=TRUE on rows to export.")

    seen = existing_keys(output_path) if append else set()
    if append:
        print(f"Found {len(confirmed_rows)} confirmed rows; "
              f"{len(seen)} already in {Path(output_path).name}")
    else:
        print(f"Found {len(confirmed_rows)} confirmed rows (rewrite mode)")

    if files_root is None and not append:
        files_root = input("Path to folder containing the source count files: ").strip().strip('"')
    files_root = Path(files_root) if files_root else None

    mode = 'a' if append else 'w'
    written = skipped_dupe = missing = 0
    by_vendor = {}

    with open(output_path, mode, encoding='utf-8') as out:
        for row in confirmed_rows:
            fname = row['file_name']
            key = dedupe_key(fname, row.get('intersection'))
            if append and key in seen:
                skipped_dupe += 1
                continue

            content = _load_file_content(files_root, fname)
            if content is None:
                print(f"  ✗ Not found: {fname}")
                missing += 1
                continue

            extraction = row_to_extraction(row)
            example = build_training_example(content, fname, extraction)
            out.write(json.dumps(example) + '\n')
            written += 1
            seen.add(key)   # guard against intra-run duplicates too
            v = row.get('vendor') or 'Unknown'
            by_vendor[v] = by_vendor.get(v, 0) + 1

    verb = "Appended" if append else "Wrote"
    print(f"\n✓ {verb} {written} training examples to {output_path}")
    if by_vendor:
        print("  By vendor: " + ", ".join(f"{k}={n}" for k, n in sorted(by_vendor.items())))
    if append and skipped_dupe:
        print(f"  Skipped {skipped_dupe} already-present (dedupe)")
    if missing:
        print(f"⚠  {missing} source files missing — skipped")
    return {"written": written, "skipped_dupe": skipped_dupe,
            "missing": missing, "by_vendor": by_vendor}


def main():
    p = argparse.ArgumentParser(description="Export confirmed labels to training JSONL")
    p.add_argument('labels', nargs='?', default="labels_template.xlsx")
    p.add_argument('output', nargs='?', default="training_data.jsonl")
    p.add_argument('--append-confirmed', action='store_true', dest='append',
                   help="Append only new confirmed rows (dedupe by file_name+intersection)")
    p.add_argument('--files-root', help="Folder containing the source count files")
    args = p.parse_args()

    if not Path(args.labels).exists():
        sys.exit(f"Labels file not found: {args.labels}")

    export(args.labels, args.output, files_root=args.files_root, append=args.append)


if __name__ == "__main__":
    main()
