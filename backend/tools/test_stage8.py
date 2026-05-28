"""
Stage 8 tests — labeling-assist harness + feedback pipeline.

Runs WITHOUT the Claude API: extraction/verification are represented by
plain dicts (the same shapes count_agent.extract / count_verifier.verify
return), so build_row / review_flag / export are tested as pure logic.

Run:
    cd backend/tools
    PYTHONPATH=.. python3.12 test_stage8.py
"""
import os, sys, json, tempfile
from pathlib import Path

# Make count_agent (in backend/) and the tools modules importable.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))   # backend/
sys.path.insert(0, str(Path(__file__).resolve().parent))          # backend/tools/

from openpyxl import Workbook, load_workbook

import populate_labels as P
import export_for_finetuning as E


# ── helpers ────────────────────────────────────────────────────────────
def make_labels_xlsx(path, rows):
    """Write a minimal Labels workbook. `rows` is a list of dicts keyed by
    COLUMNS names (+ optional 'confirmed' bool)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Labels'
    for ci, col in enumerate(P.COLUMNS, 1):
        ws.cell(row=1, column=ci, value=col)
    for ri, row in enumerate(rows, 2):
        for ci, col in enumerate(P.COLUMNS, 1):
            ws.cell(row=ri, column=ci, value=row.get(col))
    wb.save(path)


def confirmed_row(file_name, vendor, intersection, confirmed, ebl_am=10):
    return {
        'file_name': file_name, 'vendor': vendor, 'intersection': intersection,
        'AM_peak': '07:30', 'PM_peak': '16:30', 'EBL_AM': ebl_am,
        'PHF_AM': 0.9, 'PHF_PM': 0.92, 'HV_AM_pct': 2.0, 'HV_PM_pct': 1.5,
        'HV_method': 'x', 'verification_status': 'SKIPPED',
        'confirmed': confirmed, 'notes': '', 'extracted_at': '2026-01-01 00:00',
    }


# ── tests ──────────────────────────────────────────────────────────────
def test_prefill_confirmed_false():
    """Harness build_row sets confirmed=False on every row — incl. the
    empty/error path."""
    ext = {'vendor_detected': 'GRAM', 'intersections': [
        {'name': 'A & B', 'volumes': {'AM': {'EBL': 5}, 'PM': {}},
         'phf': {'AM': 0.9}, 'hv_pct': {}, 'notes': ''}]}
    ver = {'status': 'SKIPPED', 'reason': 'PDF'}
    row = P.build_row('x.pdf', ext, ver)
    assert row['confirmed'] is False, row['confirmed']

    # No-intersection path
    empty = P.build_row('y.pdf', {'vendor_detected': 'GRAM', 'intersections': []}, {})
    assert empty['confirmed'] is False

    # And the written cell is False after append_row round-trips through xlsx
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / 'out.xlsx'
        wb = Workbook(); ws = wb.active; ws.title = 'Labels'
        for ci, col in enumerate(P.COLUMNS, 1):
            ws.cell(row=1, column=ci, value=col)
        rc = P._ensure_review_header(ws)
        P.append_row(ws, row, rc)
        wb.save(out)
        wb2 = load_workbook(out); ws2 = wb2['Labels']
        conf_idx = P.COLUMNS.index('confirmed') + 1
        assert ws2.cell(row=2, column=conf_idx).value is False
    print("  test_prefill_confirmed_false: PASS")


def test_verifier_flags_surface():
    """A row failing a verifier check gets a non-empty REVIEW flag; a
    cleanly-verified row gets none."""
    mismatch_ver = {'status': 'OK', 'intersections': [{'overall': 'MISMATCH'}]}
    ext = {'vendor_detected': 'NDS', 'intersections': [
        {'name': 'A & B', 'volumes': {'AM': {}, 'PM': {}}, 'phf': {}, 'hv_pct': {}, 'notes': ''}]}
    row = P.build_row('m.xlsx', ext, mismatch_ver)
    assert row['verification_status'] == 'MISMATCH'
    assert row['review'] == 'FAIL', row['review']

    verified_ver = {'status': 'OK', 'intersections': [{'overall': 'VERIFIED'}]}
    row_ok = P.build_row('v.xlsx', ext, verified_ver)
    assert row_ok['review'] == '', row_ok['review']

    check_ver = {'status': 'OK', 'intersections': [{'overall': 'CHECK'}]}
    assert P.build_row('c.xlsx', ext, check_ver)['review'] == 'CHECK'

    # direct unit on the pure helper
    assert P.review_flag('MISMATCH') == 'FAIL'
    assert P.review_flag('VERIFIED') == ''
    assert P.review_flag('SKIPPED') == 'UNVERIFIED'
    print("  test_verifier_flags_surface: PASS")


def test_export_only_confirmed():
    """export() includes confirmed==True rows only; confirmed==False excluded."""
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        # dummy source files so _load_file_content can read content
        for fn in ('a.csv', 'b.csv', 'c.csv'):
            (d / fn).write_text("time,EBL\n07:30,10\n")
        labels = d / 'labels.xlsx'
        make_labels_xlsx(labels, [
            confirmed_row('a.csv', 'GRAM', 'A & B', True),
            confirmed_row('b.csv', 'GRAM', 'C & D', True),
            confirmed_row('c.csv', 'GRAM', 'E & F', False),   # excluded
        ])
        out = d / 'train.jsonl'
        stats = E.export(str(labels), str(out), files_root=str(d), append=False)
        assert stats['written'] == 2, stats
        lines = [json.loads(l) for l in out.read_text().splitlines() if l.strip()]
        assert len(lines) == 2
        # the unconfirmed file must not appear
        names = []
        for ex in lines:
            u = next(m['content'] for m in ex['messages'] if m['role'] == 'user')
            names.append(u.split('\n', 1)[0])
        assert not any('c.csv' in n for n in names), names
        assert any('a.csv' in n for n in names)
    print("  test_export_only_confirmed: PASS")


def test_export_dedupe():
    """Running --append-confirmed twice adds each example exactly once."""
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        for fn in ('a.csv', 'b.csv'):
            (d / fn).write_text("time,EBL\n07:30,10\n")
        labels = d / 'labels.xlsx'
        make_labels_xlsx(labels, [
            confirmed_row('a.csv', 'GRAM', 'A & B', True),
            confirmed_row('b.csv', 'GRAM', 'C & D', True),
        ])
        out = d / 'train.jsonl'

        s1 = E.export(str(labels), str(out), files_root=str(d), append=True)
        assert s1['written'] == 2, s1
        n_after_first = len([l for l in out.read_text().splitlines() if l.strip()])
        assert n_after_first == 2

        s2 = E.export(str(labels), str(out), files_root=str(d), append=True)
        assert s2['written'] == 0, s2
        assert s2['skipped_dupe'] == 2, s2
        n_after_second = len([l for l in out.read_text().splitlines() if l.strip()])
        assert n_after_second == 2, f"dedupe failed: {n_after_second} lines"
    print("  test_export_dedupe: PASS")


if __name__ == '__main__':
    tests = [
        test_prefill_confirmed_false,
        test_verifier_flags_surface,
        test_export_only_confirmed,
        test_export_dedupe,
    ]
    passed = failed = 0
    for t in tests:
        print(f"Running {t.__name__}...")
        try:
            t(); passed += 1
        except AssertionError as e:
            print(f"  FAIL: {e}"); failed += 1
        except Exception:
            import traceback; traceback.print_exc(); failed += 1
    print(f"\n{'='*50}\n  {passed} passed, {failed} failed")
    sys.exit(0 if failed == 0 else 1)
