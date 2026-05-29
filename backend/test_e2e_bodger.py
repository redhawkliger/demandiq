"""End-to-end integration test — full Bodger corridor pipeline.

Exercises the REAL functions (no mocks) across balance → link conservation →
movements review → fractional QC → Excel export → diagram render → override
propagation. This is the regression guard: if a future change silently breaks
the chain, this fails.

Run: cd backend && python3.12 test_e2e_bodger.py
"""
import sys, os, io
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook

from balancer import (
    balance_nchrp,
    fractional_volume_flags,
    compute_link_conservation,
)
from excel_export import build_workbook
from line_diagram import render_single_diagram, render_diagram_grid, FIG_INCHES, DPI

# ── Bodger 9-intersection fixture (copied from test_balancer_nchrp.py) ──────
FULL = {'EBL': 0, 'EBT': 0, 'EBR': 0, 'WBL': 0, 'WBT': 0, 'WBR': 0,
        'NBL': 0, 'NBT': 0, 'NBR': 0, 'SBL': 0, 'SBT': 0, 'SBR': 0,
        'WBU': 0, 'EBU': 0, 'NBU': 0, 'SBU': 0}


def mk(d):
    m = dict(FULL)
    m.update(d)
    return m


def bodger_fixture():
    coords = {n: {'lat': 0.0, 'lng': float(n)} for n in range(1, 10)}
    coords[6] = {'lat': 0.005, 'lng': 5.0}
    coords[5] = {'lat': 0.0, 'lng': 5.0}
    raw = {
        1: mk({'EBL': 80, 'EBT': 1000, 'EBR': 50, 'WBL': 50, 'WBT': 900, 'WBR': 40, 'NBR': 25, 'SBL': 25}),
        2: mk({'EBL': 60, 'EBT': 1050, 'EBR': 40, 'WBL': 40, 'WBT': 920, 'WBR': 50, 'NBR': 20, 'SBL': 30}),
        3: mk({'EBL': 100, 'EBT': 1020, 'EBR': 70, 'WBL': 70, 'WBT': 880, 'WBR': 60, 'NBR': 35, 'SBL': 40}),
        4: mk({'EBL': 80, 'EBT': 980, 'EBR': 60, 'WBL': 60, 'WBT': 850, 'WBR': 50, 'NBR': 30, 'SBL': 35}),
        5: mk({'EBL': 90, 'EBT': 950, 'EBR': 80, 'WBL': 80, 'WBT': 820, 'WBR': 70,
               'NBL': 30, 'NBT': 200, 'NBR': 35, 'SBL': 40, 'SBT': 180, 'SBR': 25}),
        6: mk({'NBL': 25, 'NBT': 220, 'NBR': 30, 'SBL': 35, 'SBT': 200, 'SBR': 28}),
        7: mk({'EBL': 70, 'EBT': 900, 'EBR': 50, 'WBL': 50, 'WBT': 800, 'WBR': 60, 'NBR': 25, 'SBL': 30}),
        8: mk({'EBL': 60, 'EBT': 880, 'EBR': 50, 'WBL': 50, 'WBT': 770, 'WBR': 50, 'NBR': 25, 'SBL': 25}),
        9: mk({'EBL': 50, 'EBT': 850, 'EBR': 40, 'WBL': 40, 'WBT': 750, 'WBR': 40, 'NBR': 20, 'SBL': 20}),
    }
    links = (
        [{'fromNum': a, 'toNum': b, 'direction': 'EB/WB'}
         for a, b in [(1, 2), (2, 3), (3, 4), (4, 5), (5, 7), (7, 8), (8, 9)]]
        + [{'fromNum': 5, 'toNum': 6, 'direction': 'NB/SB'}]
    )
    return raw, links, coords


STD_APPROACHES = {
    d: {'exists': True, 'lanes': [{'movement': 'LTR'}]}
    for d in ('EB', 'WB', 'NB', 'SB')
}


def build_project_state(raw, am_result, pm_result, overrides=None):
    """Assemble the frontend-shaped project_state build_workbook expects."""
    intersections = []
    for num in sorted(raw.keys()):
        vols = {mv: {'value': v} for mv, v in raw[num].items() if v}
        intersections.append({
            'num': num,
            'name': f'Int {num} & Cross St',
            'approaches': STD_APPROACHES,
            'volumes': {'AM': dict(vols), 'PM': dict(vols)},
            'balanced': {
                'AM': dict(am_result['balanced'].get(num, {})),
                'PM': dict(pm_result['balanced'].get(num, {})),
            },
            'phf': {}, 'hvPct': {},
        })
    return {
        'project': {'name': 'Bodger E2E', 'analysisPeriods': ['AM', 'PM']},
        'intersections': intersections,
        'balancing': {'overrides': overrides or {'AM': {}, 'PM': {}}},
        'analysis': {'lastRun': {'qcFlags': []}},
    }


def _geo(num, raw):
    return {
        'num': num, 'name': f'Int {num} & Cross St',
        'active_legs': ['EB', 'WB', 'NB', 'SB'], 'control': 'Signalized',
        'streets': {'NS': 'Cross St', 'EW': 'Main St'}, 'stopped_approaches': [],
    }


# ── The pipeline (run once, shared across assertions) ───────────────────────
RAW, LINKS, COORDS = bodger_fixture()
AM = balance_nchrp(RAW, LINKS, method='larger', coords=COORDS)
PM = balance_nchrp(RAW, LINKS, method='larger', coords=COORDS)   # PM mirrors AM
RESULTS = {'AM': AM, 'PM': PM}


# ── 1. BALANCE ──────────────────────────────────────────────────────────────
def test_balance():
    assert set(AM['balanced'].keys()) == set(range(1, 10)), \
        f"all 9 intersections expected; got {sorted(AM['balanced'].keys())}"
    for key in ('balanced', 'link_conservation', 'movements_review', 'qc_flags'):
        assert key in AM, f"result missing key {key}"
    print(f"  1. BALANCE: PASS (9 intersections, keys present)")


# ── 2. LINK CONSERVATION (core invariant) ───────────────────────────────────
def test_link_conservation():
    full_checked = 0
    for period in ('AM', 'PM'):
        for e in RESULTS[period]['link_conservation']:
            if e['mode'] == 'Full':
                assert abs(e['delta']) == 0, \
                    f"{period} Full link {e['fromNum']}→{e['toNum']} {e['direction']} " \
                    f"delta={e['delta']} (must be 0 post-reconciliation)"
                full_checked += 1
    assert full_checked > 0, "expected at least one Full link to check"
    print(f"  2. LINK CONSERVATION: PASS ({full_checked} Full-link deltas == 0 across AM+PM)")


# ── 3. MOVEMENTS REVIEW ─────────────────────────────────────────────────────
def test_movements_review():
    review = {'AM': AM['movements_review'], 'PM': PM['movements_review']}
    for period in ('AM', 'PM'):
        assert isinstance(review[period], list), f"{period} movements_review not a list"
        assert len(review[period]) > 0, f"{period} movements_review empty"
        for e in review[period]:
            assert set(e.keys()) == {'int_num', 'movement', 'raw', 'balanced', 'delta'}, e.keys()
            assert e['delta'] == e['balanced'] - e['raw'], \
                f"delta mismatch: {e}"
    print(f"  3. MOVEMENTS REVIEW: PASS "
          f"(AM {len(review['AM'])}, PM {len(review['PM'])}; delta == bal-raw)")


# ── 4. FRACTIONAL QC ────────────────────────────────────────────────────────
def test_fractional_qc():
    # Clean integer raw → zero fractional flags
    assert fractional_volume_flags(RAW, period='AM') == [], "clean raw should fire 0 flags"

    # One deliberately-fractional value
    bad = {n: dict(v) for n, v in RAW.items()}
    bad[3]['WBT'] = 102.5
    flags = fractional_volume_flags(bad, period='AM')
    assert len(flags) == 1, f"expected exactly 1 fractional flag, got {len(flags)}"
    f = flags[0]
    assert f['check'] == 'FRACTIONAL-VOLUME' and f['int'] == 3
    assert 'WBT' in f['message'] and '102.5' in f['message'], f['message']
    print(f"  4. FRACTIONAL QC: PASS (1 flag on Int3 WBT 102.5, 0 on clean raw)")


# ── 5. EXCEL EXPORT ─────────────────────────────────────────────────────────
def test_excel_export():
    state = build_project_state(RAW, AM, PM)
    xlsx = build_workbook(state)
    assert isinstance(xlsx, (bytes, bytearray)) and len(xlsx) > 0, "workbook bytes empty"

    wb = load_workbook(io.BytesIO(xlsx), data_only=True)
    # PRD-02: workbook is 3 tabs — Provenance / Combined View / Raw Movements.
    assert wb.sheetnames == ['Provenance', 'Combined View', 'Raw Movements'], wb.sheetnames

    # Spot-check a Combined View row: AM Δ == AM Balanced − AM Raw
    ws = wb['Combined View']
    headers = [ws.cell(row=1, column=c).value for c in range(1, 8)]
    assert headers == ['Movement', 'AM Raw', 'AM Balanced', 'AM Δ', 'PM Raw', 'PM Balanced', 'PM Δ'], headers
    checked = 0
    for r in range(2, ws.max_row + 1):
        mv = ws.cell(row=r, column=1).value
        if mv in (None, 'TOTAL') or not isinstance(mv, str) or len(mv) != 3:
            continue
        raw_v = ws.cell(row=r, column=2).value
        bal_v = ws.cell(row=r, column=3).value
        dlt_v = ws.cell(row=r, column=4).value
        # skip override-marked cells (string with "*") — covered in test 7
        if isinstance(bal_v, str) or raw_v is None or bal_v is None:
            continue
        assert dlt_v == bal_v - raw_v, f"row {r} {mv}: Δ {dlt_v} != {bal_v}-{raw_v}"
        checked += 1
    assert checked > 0, "expected at least one numeric Combined View row to spot-check"
    print(f"  5. EXCEL EXPORT: PASS (sheets incl. Combined View; {checked} Δ rows verified)")


# ── 6. DIAGRAM RENDER ───────────────────────────────────────────────────────
def test_diagram_render():
    # Single
    u = {mv: v for mv, v in RAW[1].items() if v}
    png = render_single_diagram(_geo(1, RAW), u, AM['balanced'][1], 'combined')
    assert png[:4] == b'\x89PNG', "single diagram is not a PNG"
    assert len(png) > 1000, "single diagram suspiciously small"

    # Grid of all 9 → 3 wide
    cells = [{
        'geo': _geo(n, RAW),
        'u_vols': {mv: v for mv, v in RAW[n].items() if v},
        'b_vols': AM['balanced'][n],
        'overrides': {},
        'period_label': 'AM',
    } for n in range(1, 10)]
    grid = render_diagram_grid(cells, 'combined', cols=3)
    assert grid[:4] == b'\x89PNG', "grid is not a PNG"
    from PIL import Image
    w, h = Image.open(io.BytesIO(grid)).size
    cell_px = int(FIG_INCHES * DPI)
    assert w == 3 * cell_px, f"grid width {w} != 3 cells"
    assert h == 3 * cell_px, f"grid height {h} != 3 rows (9 cells / 3 cols)"
    print(f"  6. DIAGRAM RENDER: PASS (single PNG + 9-cell grid {w}x{h})")


# ── 7. OVERRIDE PROPAGATION ─────────────────────────────────────────────────
def test_override_propagation():
    bal_ebl = int(AM['balanced'][4].get('EBL', 0))
    raw_ebl = int(RAW[4]['EBL'])
    # Override to raw value; guarantee it differs from balanced so the
    # conservation break is deterministic.
    ovr_val = raw_ebl if raw_ebl != bal_ebl else bal_ebl + 50
    overrides = {'AM': {'4': {'EBL': ovr_val}}, 'PM': {}}

    # (a) Excel Combined View cell carries the "*" marker + override value
    state = build_project_state(RAW, AM, PM, overrides=overrides)
    wb = load_workbook(io.BytesIO(build_workbook(state)), data_only=True)
    ws = wb['Combined View']
    # find the Int 4 section, then its EBL row
    in_int4 = False
    ebl_cell = None
    for r in range(2, ws.max_row + 1):
        first = ws.cell(row=r, column=1).value
        if isinstance(first, str) and first.startswith('Int 4 '):
            in_int4 = True
            continue
        if isinstance(first, str) and first.startswith('Int ') and not first.startswith('Int 4 '):
            in_int4 = False
        if in_int4 and first == 'EBL':
            ebl_cell = ws.cell(row=r, column=3).value   # AM Balanced
            break
    assert ebl_cell == f"{ovr_val}*", f"expected '{ovr_val}*' in Int4 EBL AM cell; got {ebl_cell!r}"

    # (b) The EB link into Int 4 now shows a nonzero conservation delta
    balanced_ovr = {n: dict(v) for n, v in AM['balanced'].items()}
    balanced_ovr[4]['EBL'] = ovr_val
    lc = compute_link_conservation(balanced_ovr, AM['resolved_links'])
    entry = next((e for e in lc if e['toNum'] == 4 and e['direction'] == 'EB'), None)
    assert entry is not None, "no (→4, EB) link found"
    assert abs(entry['delta']) > 0, \
        f"override should break conservation on (3→4,EB); delta={entry['delta']}"
    print(f"  7. OVERRIDE PROPAGATION: PASS "
          f"(Int4 EBL→{ovr_val}* in Excel; (3→4,EB) delta now {entry['delta']:+.0f})")


if __name__ == '__main__':
    tests = [
        test_balance,
        test_link_conservation,
        test_movements_review,
        test_fractional_qc,
        test_excel_export,
        test_diagram_render,
        test_override_propagation,
    ]
    passed = failed = 0
    print("\nBodger end-to-end pipeline:")
    for t in tests:
        try:
            t(); passed += 1
        except AssertionError as e:
            print(f"  FAIL {t.__name__}: {e}"); failed += 1
        except Exception:
            import traceback; traceback.print_exc(); failed += 1
    print(f"\n{'='*54}\n  {passed} passed, {failed} failed")
    sys.exit(0 if failed == 0 else 1)
