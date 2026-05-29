"""PRD-02 — Excel export honors engineer overrides + 3-tab schema.

Verifies that project_state.balancing.overrides is reflected in:
  1. Combined View — balanced cell carries "*" suffix on overridden movements.
  2. Raw Movements — the dedicated "AM Override" / "PM Override" column
     carries the engineer value.

Also covers the Provenance tab structure (header block + per-intersection
convergence table + QC flags section) per PRD-02 AC.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import io
from openpyxl import load_workbook
from excel_export import build_workbook


def test_export_honors_overrides():
    # Minimal project state with one intersection + one override
    project_state = {
        'project': {'name': 'Test', 'client': '', 'projectNumber': '', 'analyst': '',
                    'checker': '', 'date': '', 'scenario': '', 'projectType': '',
                    'analysisPeriods': ['AM']},
        'intersections': [{
            'num': 5,
            'name': 'Main St & 1st Ave',
            'approaches': {
                'EB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
                'WB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
                'NB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
                'SB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
            },
            'volumes': {'AM': {
                'EBL': {'value': 90, '_source': 'agent'},
                'EBT': {'value': 950, '_source': 'agent'},
                'EBR': {'value': 80, '_source': 'agent'},
            }},
            'balanced': {'AM': {
                'EBL': 95, 'EBT': 1050, 'EBR': 85,
            }},
        }],
        'balancing': {
            'overrides': {
                'AM': {'5': {'EBT': 950}},   # engineer clamped EBT back to raw
                'PM': {},
            },
        },
        'analysis': {'lastRun': {'qcFlags': []}},
    }

    xlsx_bytes = build_workbook(project_state)
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)

    # ── Sheet schema ───────────────────────────────────────────────────
    assert wb.sheetnames == ['Provenance', 'Combined View', 'Raw Movements'], wb.sheetnames

    # ── Combined View carries the override "*" marker on EBT balanced ──
    ws = wb['Combined View']
    ebt_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == 'EBT':
            ebt_row = r
            break
    assert ebt_row, 'EBT row not found in Combined View'
    am_balanced = ws.cell(row=ebt_row, column=3).value
    assert am_balanced == '950*', f"expected '950*' in EBT balanced; got {am_balanced!r}"
    print(f"  Combined View EBT AM balanced shows override marker: {am_balanced!r}")

    # ── Raw Movements carries the override in the AM Override column ───
    raw_ws = wb['Raw Movements']
    expected_headers = [
        'Int #', 'Movement',
        'AM Raw', 'AM Balanced', 'AM Δ', 'AM Override',
        'PM Raw', 'PM Balanced', 'PM Δ', 'PM Override',
    ]
    for col, label in enumerate(expected_headers, start=1):
        assert raw_ws.cell(row=1, column=col).value == label, \
            f"header {col}: got {raw_ws.cell(row=1, column=col).value!r}, expected {label!r}"
    # Find the EBT row in Raw Movements (Int #5, Movement EBT)
    ebt_raw_row = None
    for r in range(2, raw_ws.max_row + 1):
        if raw_ws.cell(row=r, column=1).value == 5 and raw_ws.cell(row=r, column=2).value == 'EBT':
            ebt_raw_row = r
            break
    assert ebt_raw_row, 'EBT row not found in Raw Movements'
    am_raw      = raw_ws.cell(row=ebt_raw_row, column=3).value
    am_balanced = raw_ws.cell(row=ebt_raw_row, column=4).value
    am_override = raw_ws.cell(row=ebt_raw_row, column=6).value
    assert am_raw == 950, f"AM Raw: {am_raw}"
    assert am_balanced == 1050, f"AM Balanced (pre-override): {am_balanced}"
    assert am_override == 950, f"AM Override: {am_override}"
    print(f"  Raw Movements EBT row: raw={am_raw} bal={am_balanced} override={am_override}")
    print(f"  test_export_honors_overrides: PASS")


def test_export_clean_when_no_overrides():
    """No overrides — Raw Movements emits the table, Override columns blank."""
    project_state = {
        'project': {'name': 'Test', 'analysisPeriods': []},
        'intersections': [{
            'num': 1, 'name': 'Foo',
            'approaches': {'EB': {'exists': True, 'lanes': [{'movement': 'LTR'}]}},
            'volumes': {'AM': {'EBT': {'value': 100, '_source': 'agent'}}},
            'balanced': {'AM': {'EBT': 100}},
        }],
        'balancing': {'overrides': {'AM': {}, 'PM': {}}},
        'analysis': {'lastRun': {}},
    }
    xlsx_bytes = build_workbook(project_state)
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    assert wb.sheetnames == ['Provenance', 'Combined View', 'Raw Movements'], wb.sheetnames

    raw_ws = wb['Raw Movements']
    # EBT row, AM Override (col 6) should be None (blank)
    found = False
    for r in range(2, raw_ws.max_row + 1):
        if raw_ws.cell(row=r, column=2).value == 'EBT':
            assert raw_ws.cell(row=r, column=6).value is None, \
                f"AM Override should be blank; got {raw_ws.cell(row=r, column=6).value!r}"
            found = True
            break
    assert found, 'EBT row not found in Raw Movements'
    print(f"  test_export_clean_when_no_overrides: PASS")


def test_combined_view_sheet():
    """Stage 10.2 — Combined View sheet: 7 columns, Δ = Bal − Raw, and
    override balanced cells carry the "*" marker. Preserved unchanged."""
    project_state = {
        'project': {'name': 'Test', 'analysisPeriods': ['AM', 'PM']},
        'intersections': [{
            'num': 3,
            'name': 'Main St & 1st Ave',
            'approaches': {
                'EB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
                'WB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
                'NB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
                'SB': {'exists': True, 'lanes': [{'movement': 'LTR'}]},
            },
            'volumes': {
                'AM': {'EBT': {'value': 1000}, 'EBL': {'value': 80}},
                'PM': {'EBT': {'value': 1200}},
            },
            'balanced': {
                'AM': {'EBT': 1050, 'EBL': 60},   # EBL dropped 80→60
                'PM': {'EBT': 1200},
            },
        }],
        'balancing': {'overrides': {'AM': {'3': {'EBT': 1000}}, 'PM': {}}},
        'analysis': {'lastRun': {'qcFlags': []}},
    }

    xlsx_bytes = build_workbook(project_state)
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)

    assert 'Combined View' in wb.sheetnames, wb.sheetnames
    ws = wb['Combined View']

    expected = ['Movement', 'AM Raw', 'AM Balanced', 'AM Δ', 'PM Raw', 'PM Balanced', 'PM Δ']
    actual = [ws.cell(row=1, column=c).value for c in range(1, 8)]
    assert actual == expected, f"header mismatch: {actual}"

    ebl_row = None
    ebt_row = None
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v == 'EBL':
            ebl_row = r
        elif v == 'EBT':
            ebt_row = r
    assert ebl_row, "EBL row not found"
    am_raw = ws.cell(row=ebl_row, column=2).value
    am_bal = ws.cell(row=ebl_row, column=3).value
    am_delta = ws.cell(row=ebl_row, column=4).value
    assert am_raw == 80, am_raw
    assert am_bal == 60, am_bal
    assert am_delta == am_bal - am_raw == -20, am_delta
    print(f"  EBL row: raw={am_raw} bal={am_bal} Δ={am_delta}")

    ebt_am_bal = ws.cell(row=ebt_row, column=3).value
    assert ebt_am_bal == '1000*', f"expected '1000*' override marker; got {ebt_am_bal!r}"
    print(f"  EBT AM override marker: {ebt_am_bal!r}")

    totals = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert 'TOTAL' in totals, "TOTAL subtotal row missing"
    print(f"  test_combined_view_sheet: PASS")


def test_provenance_tab_structure():
    """PRD-02 AC: Provenance tab has the header block, per-intersection
    convergence table, and QC flags section populated from lastRun.config +
    intConverged + iterations + finalResidual + qcFlags."""
    project_state = {
        'project': {'name': 'PRD-02 fixture', 'analysisPeriods': ['AM', 'PM']},
        'intersections': [
            {'num': 1, 'name': 'A', 'approaches': {},
             'volumes': {'AM': {}, 'PM': {}}, 'balanced': {'AM': {}, 'PM': {}}},
            {'num': 2, 'name': 'B', 'approaches': {},
             'volumes': {'AM': {}, 'PM': {}}, 'balanced': {'AM': {}, 'PM': {}}},
        ],
        'balancing': {
            'linkResolution': 'larger',
            'corridorDir': 'NB/SB',
            'peakMode': 'Corridor',
            'amWindow': {'start': '07:00', 'end': '09:00'},
            'pmWindow': {'start': '16:00', 'end': '18:00'},
            'convergenceThreshold': 0.5,
            'maxIterations': 500,
            'overrides': {'AM': {}, 'PM': {}},
        },
        'analysis': {'lastRun': {
            'timestamp': '2026-05-29T19:00:00.000Z',
            'method': 'larger',
            'iterations': {'AM': {'1': 87, '2': 114}, 'PM': {'1': 92, '2': 121}},
            'intConverged': {'AM': {'1': True, '2': True}, 'PM': {'1': True, '2': False}},
            'finalResidual': {'AM': {'1': 0.21, '2': 0.34}, 'PM': {'1': 0.27, '2': 0.41}},
            'qcFlags': [
                {'check': 'NCHRP-LINK-DELTA', 'period': 'PM', 'message': 'Link 1→2 delta 1.2 veh'},
            ],
            'config': {
                'linkResolution': 'average', 'peakMode': 'Independent',
                'amWindow': {'start': '06:30', 'end': '08:30'},
                'pmWindow': {'start': '16:30', 'end': '18:30'},
                'convergenceThreshold': 0.25, 'maxIterations': 750,
                'corridorDir': 'EB/WB',
            },
        }},
    }

    wb = load_workbook(io.BytesIO(build_workbook(project_state)), data_only=True)
    assert wb.sheetnames[0] == 'Provenance', wb.sheetnames
    ws = wb['Provenance']

    # Build a key→value map of the header block (cols A,B; first 10 rows)
    header = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 11)}
    assert header['Project']        == 'PRD-02 fixture',           header['Project']
    assert header['Run at']         == '2026-05-29T19:00:00.000Z', header['Run at']
    # Method/Direction/Peak come from lastRun.config (NOT current balancing)
    assert header['Method']         == 'Average',                  header['Method']
    assert header['Direction']      == 'EB/WB',                    header['Direction']
    assert header['Peak mode']      == 'Independent',              header['Peak mode']
    assert header['AM window']      == '06:30 – 08:30',            header['AM window']
    assert header['PM window']      == '16:30 – 18:30',            header['PM window']
    assert header['Tolerance']      == '0.25 veh',                 header['Tolerance']
    assert header['Max iterations'] == 750,                        header['Max iterations']
    assert header['Intersections']  == 2,                          header['Intersections']
    print(f"  Provenance header block reads from lastRun.config (not current balancing).")

    # Locate the per-intersection convergence section
    conv_header_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == 'Per-intersection convergence':
            conv_header_row = r
            break
    assert conv_header_row, 'convergence section header not found'
    col_headers = [ws.cell(row=conv_header_row + 1, column=c).value for c in range(1, 9)]
    assert col_headers == [
        'Int #', 'Name', 'AM iter', 'AM conv', 'AM final residual',
        'PM iter', 'PM conv', 'PM final residual',
    ], col_headers
    # Int #1 row
    r1 = conv_header_row + 2
    assert ws.cell(row=r1, column=1).value == 1
    assert ws.cell(row=r1, column=3).value == 87        # AM iter
    assert ws.cell(row=r1, column=4).value == 'Y'        # AM conv
    assert ws.cell(row=r1, column=5).value == 0.21       # AM final residual
    assert ws.cell(row=r1, column=6).value == 92         # PM iter
    assert ws.cell(row=r1, column=7).value == 'Y'        # PM conv
    # Int #2 row — PM did not converge
    r2 = conv_header_row + 3
    assert ws.cell(row=r2, column=1).value == 2
    assert ws.cell(row=r2, column=7).value == 'N'
    print(f"  Per-intersection convergence table populates from iterations + intConverged + finalResidual.")

    # QC flags section — one row present with the seeded flag
    qc_header_row = None
    for r in range(conv_header_row, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == 'QC flags':
            qc_header_row = r
            break
    assert qc_header_row, 'QC flags section header not found'
    # row qc_header_row+1 is column headers; row qc_header_row+2 is first data
    check = ws.cell(row=qc_header_row + 2, column=1).value
    period = ws.cell(row=qc_header_row + 2, column=2).value
    msg    = ws.cell(row=qc_header_row + 2, column=3).value
    assert check == 'NCHRP-LINK-DELTA', check
    assert period == 'PM', period
    assert 'Link 1→2' in (msg or ''), msg
    print(f"  QC flags table renders the qcFlags entries.")
    print(f"  test_provenance_tab_structure: PASS")


def test_provenance_qc_empty_renders_all_clean():
    """When qcFlags is empty, the QC section shows an explicit 'All clean'
    row rather than a bare section header (PRD AC)."""
    project_state = {
        'project': {'name': 'clean run', 'analysisPeriods': ['AM']},
        'intersections': [],
        'balancing': {'overrides': {'AM': {}, 'PM': {}}},
        'analysis': {'lastRun': {
            'timestamp': '2026-05-29T20:00:00.000Z',
            'qcFlags': [],
            'config': {
                'linkResolution': 'larger', 'peakMode': 'Corridor',
                'amWindow': {'start': '07:00', 'end': '09:00'},
                'pmWindow': {'start': '16:00', 'end': '18:00'},
                'convergenceThreshold': 0.5, 'maxIterations': 500,
                'corridorDir': 'NB/SB',
            },
        }},
    }
    wb = load_workbook(io.BytesIO(build_workbook(project_state)), data_only=True)
    ws = wb['Provenance']
    qc_header_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == 'QC flags':
            qc_header_row = r
            break
    assert qc_header_row, 'QC flags section header not found'
    all_clean = ws.cell(row=qc_header_row + 2, column=1).value
    assert all_clean == 'All clean', f"expected 'All clean'; got {all_clean!r}"
    print(f"  test_provenance_qc_empty_renders_all_clean: PASS")


def test_provenance_legacy_run_degrades_gracefully():
    """Legacy (pre-NCHRP) run without lastRun.config + per-int fields renders
    '—' in the metadata block rather than crashing (PRD risk note)."""
    project_state = {
        'project': {'name': 'legacy', 'analysisPeriods': ['AM']},
        'intersections': [{'num': 1, 'name': 'A',
                           'approaches': {}, 'volumes': {}, 'balanced': {}}],
        'balancing': {'overrides': {'AM': {}, 'PM': {}}},
        # No config; no intConverged/iterations/finalResidual; legacy 'convergence'.
        'analysis': {'lastRun': {
            'timestamp': '2026-05-01T00:00:00.000Z',
            'convergence': {'AM': {'iterations_used': 12}, 'PM': {'iterations_used': 14}},
            'qcFlags': [],
        }},
    }
    # Should not raise.
    wb = load_workbook(io.BytesIO(build_workbook(project_state)), data_only=True)
    assert wb.sheetnames == ['Provenance', 'Combined View', 'Raw Movements'], wb.sheetnames
    ws = wb['Provenance']
    header = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 11)}
    # Missing fields fall back to '—'
    assert header['Method'] == '—', header['Method']
    assert header['Tolerance'] == '—', header['Tolerance']
    assert header['Max iterations'] == '—', header['Max iterations']
    print(f"  test_provenance_legacy_run_degrades_gracefully: PASS")


if __name__ == '__main__':
    tests = [
        test_export_honors_overrides,
        test_export_clean_when_no_overrides,
        test_combined_view_sheet,
        test_provenance_tab_structure,
        test_provenance_qc_empty_renders_all_clean,
        test_provenance_legacy_run_degrades_gracefully,
    ]
    passed = 0; failed = 0
    for t in tests:
        print(f"Running {t.__name__}...")
        try:
            t()
            passed += 1
        except AssertionError as e:
            print(f"  FAIL: {e}")
            failed += 1
        except Exception as e:
            import traceback
            traceback.print_exc()
            failed += 1
    print(f"\n{'='*50}\n  {passed} passed, {failed} failed")
    sys.exit(0 if failed == 0 else 1)
